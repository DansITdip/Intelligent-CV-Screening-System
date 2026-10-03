"""
CV Parser Service
-----------------
Extracts readable text from PDF and DOCX CV files.

Supported formats:
    - PDF
    - DOCX
"""

import os


def extract_text_from_pdf(file_path):
    """
    Extract text from a PDF CV.
    """

    try:
        from pypdf import PdfReader

        reader = PdfReader(file_path)

        extracted_text = []

        for page in reader.pages:
            text = page.extract_text()

            if text:
                extracted_text.append(text)

        return "\n".join(extracted_text).strip()

    except ImportError:
        raise RuntimeError(
            "PDF support is not installed. "
            "Install the 'pypdf' package."
        )

    except Exception as error:
        raise RuntimeError(
            f"Unable to read PDF file: {error}"
        )


def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX CV.
    """

    try:
        from docx import Document

        document = Document(file_path)

        extracted_text = []

        # Extract normal paragraphs
        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                extracted_text.append(text)

        # Extract text from tables
        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    extracted_text.append(" | ".join(row_text))

        return "\n".join(extracted_text).strip()

    except ImportError:
        raise RuntimeError(
            "DOCX support is not installed. "
            "Install the 'python-docx' package."
        )

    except Exception as error:
        raise RuntimeError(
            f"Unable to read DOCX file: {error}"
        )


def extract_cv_text(file_path):
    """
    Automatically detect the CV file type
    and extract its text.
    """

    if not file_path:
        raise ValueError("No CV file path was provided.")

    if not os.path.isfile(file_path):
        raise FileNotFoundError(
            f"CV file not found: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(file_path)

    elif extension == ".docx":

        return extract_text_from_docx(file_path)

    else:

        raise ValueError(
            "Unsupported CV format. "
            "Only PDF and DOCX files are supported."
        )


def validate_cv_file(file_path, max_size_mb=10):
    """
    Validate that a CV exists, has a supported extension,
    and does not exceed the maximum file size.
    """

    if not file_path:
        return False, "No file was provided."

    if not os.path.isfile(file_path):
        return False, "CV file does not exist."

    extension = os.path.splitext(file_path)[1].lower()

    allowed_extensions = {".pdf", ".docx"}

    if extension not in allowed_extensions:
        return False, "Only PDF and DOCX files are allowed."

    max_size_bytes = max_size_mb * 1024 * 1024

    file_size = os.path.getsize(file_path)

    if file_size > max_size_bytes:
        return False, (
            f"File is too large. Maximum allowed size is "
            f"{max_size_mb} MB."
        )

    if file_size == 0:
        return False, "The uploaded CV is empty."

    return True, "CV file is valid."


def get_cv_file_info(file_path):
    """
    Return basic information about a CV file.
    """

    if not os.path.isfile(file_path):
        raise FileNotFoundError(
            f"CV file not found: {file_path}"
        )

    filename = os.path.basename(file_path)

    extension = os.path.splitext(filename)[1].lower()

    file_size = os.path.getsize(file_path)

    file_size_kb = round(file_size / 1024, 2)

    return {
        "filename": filename,
        "extension": extension,
        "size_bytes": file_size,
        "size_kb": file_size_kb
    }


def clean_extracted_text(text):
    """
    Clean extracted CV text by removing unnecessary
    spaces and blank lines.
    """

    if not text:
        return ""

    lines = []

    for line in text.splitlines():

        cleaned_line = " ".join(line.split())

        if cleaned_line:
            lines.append(cleaned_line)

    return "\n".join(lines)


def parse_cv(file_path):
    """
    Complete CV parsing pipeline.

    Returns:
        Dictionary containing:
            - filename
            - file type
            - file size
            - extracted text
            - character count
            - word count
    """

    valid, message = validate_cv_file(file_path)

    if not valid:
        raise ValueError(message)

    file_info = get_cv_file_info(file_path)

    raw_text = extract_cv_text(file_path)

    cleaned_text = clean_extracted_text(raw_text)

    if not cleaned_text:
        raise ValueError(
            "No readable text was found in the CV."
        )

    return {
        "filename": file_info["filename"],
        "file_type": file_info["extension"].replace(".", "").upper(),
        "size_bytes": file_info["size_bytes"],
        "size_kb": file_info["size_kb"],
        "text": cleaned_text,
        "character_count": len(cleaned_text),
        "word_count": len(cleaned_text.split())
    }


if __name__ == "__main__":

    print("=" * 60)
    print("INTELLIGENT CV SCREENING SYSTEM")
    print("CV PARSER SERVICE")
    print("=" * 60)

    print()
    print("Supported CV formats:")
    print("  ✓ PDF")
    print("  ✓ DOCX")
    print()
    print("Maximum file size: 10 MB")
    print()
    print("CV parser service is ready.")
    print("=" * 60)